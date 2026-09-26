"""Read-only, bounded label inventory of the 21 original CETPRO workbooks."""
from pathlib import Path
import re
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1] / "sources/templates/originals/xlsx"
words = re.compile(r"cetpro|instituc|dre|modular|modulo|m[oó]dulo|programa|periodo|per[ií]odo|estudiant|alumno|apellidos|nombres|documento|dni|gesti[oó]n|resoluci[oó]n|departamento|provincia|distrito|unidad|cr[eé]dito|hora|nivel|plan|asistencia|evaluaci[oó]n|indicador|efsrt|acta|certificado|t[ií]tulo|secci[oó]n|grupo|director|direcci[oó]n|tel[eé]fono|fecha|empresa|nota|capacidad", re.I)
for file in sorted(root.glob("*.xlsx")):
    print(f"\n{file.name}")
    book = load_workbook(file, read_only=True, data_only=True)
    for sheet in book.worksheets:
        print(f"  HOJA {sheet.title}")
        found = []
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str):
                    continue
                value = " ".join(value.split())
                if 2 <= len(value) <= 110 and words.search(value):
                    found.append(f"{cell.coordinate}={value}")
        print("    " + " | ".join(found[:35]))
    book.close()
